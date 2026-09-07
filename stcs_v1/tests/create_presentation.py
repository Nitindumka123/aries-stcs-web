import os
try:
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.enum.text import PP_ALIGN
    from pptx.dml.color import RGBColor
except ImportError:
    print("Please install python-pptx: pip install python-pptx")
    exit()

def create_presentation():
    prs = Presentation()

    def set_background(slide, r, g, b):
        background = slide.background
        fill = background.fill
        fill.solid()
        fill.fore_color.rgb = RGBColor(r, g, b)

    def add_slide(title_text, bullet_points):
        slide_layout = prs.slide_layouts[1]
        slide = prs.slides.add_slide(slide_layout)
        
        # Title styling
        title = slide.shapes.title
        title.text = title_text
        
        # Content styling
        content = slide.placeholders[1]
        tf = content.text_frame
        tf.word_wrap = True
        
        for point in bullet_points:
            p = tf.add_paragraph()
            p.text = point
            p.level = 0
            p.space_after = Pt(10)

    # --- Slide 1: Title Slide ---
    title_slide_layout = prs.slide_layouts[0]
    slide = prs.slides.add_slide(title_slide_layout)
    title = slide.shapes.title
    subtitle = slide.placeholders[1]
    
    title.text = "Modernization of the 104cm Sampurnanand Telescope Control System"
    subtitle.text = "Intelligent Motion Control, Safety Interlocks, and Stellarium Integration\nBeluwakhan Observatory"

    # --- Slide 2: System Overview ---
    add_slide("System Overview", [
        "Control Objective: Modernize the 2-pier English Equatorial Mount for automated slewing and tracking.",
        "Interface: A unified Master Control Program (MCP) acting as a bridge between astronomical software and hardware.",
        "Hardware: Arduino Mega 2560 controlling high-torque relay modules for RA/Dec motions.",
        "Software Stack: PyQt6 for GUI, Astropy for high-precision astronomy math, and Stellarium for visual commanding."
    ])

    # --- Slide 3: Communication Protocol & Reliability ---
    add_slide("Communication & Reliability", [
        "Binary-to-ASCII Bypass: Implementation of an 0x30 offset to prevent CH340 serial drivers from dropping null bytes.",
        "Heartbeat Watchdog: Hardware-level 2-second timeout on the Arduino; all relays cut if the PC signal is lost.",
        "10Hz Control Loop: Real-time relay state updates sent every 100ms for fluid motion and responsiveness.",
        "Native Stellarium Protocol: Direct binary TCP communication, removing the need for third-party middleware."
    ])

    # --- Slide 4: Intelligent Motion Control ---
    add_slide("Cascading Speed Controller", [
        "Variable Speed Simulation: Emulates multi-speed motors (Coarse, Fine 1, Fine 2) using fixed-rate relays.",
        "Dynamic Speed Selection: Software automatically chooses the highest safe speed based on the remaining angular error.",
        "Automatic Deceleration: System 'steps down' through speeds as the target is approached to prevent mechanical overshoot.",
        "True Sidereal Tracking: Tracking rate calculated at 15.041°/hour to compensate for Earth's rotation vs Solar time."
    ])

    # --- Slide 5: Mechanical Safety & Interlocks ---
    add_slide("Safety & Interlocks", [
        "Hardware-Level Exclusion: Arduino firmware prohibits opposing motions (East/West or North/South) regardless of PC commands.",
        "Static Horizon Limits: Hard-coded 12° altitude and Declination boundaries to protect the optical tube.",
        "Live Path Monitoring: 10Hz calculation of instantaneous Altitude; system triggers Emergency Stop if the trajectory dips below the horizon.",
        "Mechanical Protection: Prevents gear binding or dome strikes during massive cross-sky slews."
    ])

    # --- Slide 6: Intelligent Waypoint Routing ---
    add_slide("Advanced Pathfinding (Waypoint Algorithm)", [
        "The Problem: Shortest-path diagonal slews can mathematically drag the telescope 'through the ground' in the South.",
        "The Solution: Predictive path simulation before any hardware actuation occurs.",
        "Dynamic Stepping Stones: If a ground-strike is detected, the MCP generates an optimized intermediate Waypoint.",
        "Efficiency: Automatically finds the lowest safe Declination lift required, minimizing wear and slew time compared to standard Dog-Legging."
    ])

    # --- Slide 7: Mathematical Precision ---
    add_slide("Telemetry & Math Engine", [
        "Local Sidereal Time (LST): Real-time calculation using exact observatory longitude and high-precision Astropy kernels.",
        "Coordinate Transformation: Dynamic conversion between ICRS (RA/Dec) and Horizontal (Alt/Az) systems.",
        "Delta-Time Physics: The simulation loop uses a microsecond-accurate clock to prevent drift over long tracking sessions.",
        "Simulation Mode: Full operational capability without encoders, utilizing kinematic modeling for positional feedback."
    ])

    # --- Slide 8: Future Roadmap ---
    add_slide("Future Roadmap", [
        "GPS Integration: Utilizing atomic GPS timing for nanosecond-accurate FITS header timestamping and LST sync.",
        "Encoder Feedback: Transitioning from kinematic simulation to closed-loop hardware feedback.",
        "Variable Frequency Drives: Potential move from relay 'Bang-Bang' control to smooth acceleration curves.",
        "Remote Observatory Ops: Full remote capability via centralized MCP dashboard."
    ])

    # Save the file
    file_path = "Telescope_Control_System_Presentation.pptx"
    prs.save(file_path)
    print(f"Presentation saved to: {file_path}")

if __name__ == "__main__":
    create_presentation()