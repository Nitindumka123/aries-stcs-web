"""
Phase 2 workspace pages + read-only JSON APIs + safe command gate.

WHY this module exists:
  Phase 1 proved authentication. Phase 2 adds the observatory workspaces
  (CONTROL / CHECKUP / HISTORY / CAMERA / SYSTEM / ADMIN) on top of the
  same session auth, bcrypt passwords, and CSRF protection.

SAFETY MODEL (read carefully):
  - All *telemetry* endpoints here are READ-ONLY (Alpaca GETs, UDP listen,
    config files). They never move hardware.
  - All *command* intents go through POST /api/command/validate, which
    mirrors the existing STCS validation sets (VALID_COMMANDS / speeds /
    directions from telemetry_server.py) and then REFUSES execution unless
    STCS_COMMANDS_ENABLED=1 AND the caller holds the control lock. During
    development/testing commands are validated only — nothing moves.
  - The browser never touches Arduino/serial/relay code paths. A future
    reviewed adapter may bridge the validated intent to the existing
    MotionState/SlewEngine path; that bridge does not exist yet and must
    pass RULE 20 review before any physical test.
"""

import csv
import io
import json
import os
import time

from fastapi import APIRouter, HTTPException, Request, status
from fastapi.responses import HTMLResponse, JSONResponse, RedirectResponse, StreamingResponse

from . import auth
from . import observatory as obs_svc
from . import obs_store
from . import command_defs
from .routes import _csrf_token, _csrf_validate_form, template_response

router = APIRouter(tags=["observatory"])

COMMANDS_ENABLED = os.environ.get("STCS_COMMANDS_ENABLED", "0") == "1"

# Command validation sets (shared definitions)
VALID_COMMANDS = command_defs.VALID_COMMANDS
VALID_SPEEDS = command_defs.VALID_SPEEDS
VALID_DIRECTIONS = command_defs.VALID_DIRECTIONS


def _require_user(request: Request):
    user = auth.get_current_user(request)
    if user is None:
        raise HTTPException(status_code=status.HTTP_401_UNAUTHORIZED,
                            detail="Authentication required")
    return user


def _page_ctx(request: Request, user, snap, extra=None):
    safety = (snap.get("safety") or {}).get("overall", "UNKNOWN")
    alert_cls = {"SAFE": "safe", "WARNING": "warn", "CRITICAL": "crit"}.get(safety, "unk")
    alert_text = {
        "safe": "ALL SYSTEMS NOMINAL",
        "warn": "ADVISORY — CHECK SAFETY PANEL",
        "crit": "SAFETY CONDITION — REVIEW BEFORE OPERATING",
    }.get(alert_cls, "TELEMETRY UNAVAILABLE — STATE UNKNOWN")
    # Use telemetry_status from unified snapshot (prefers WS, falls back to Alpaca)
    telemetry_status = snap.get("telemetry_status", "UNKNOWN")
    led = "ok" if telemetry_status == "LIVE" else ("warn" if telemetry_status == "STALE" else "bad")
    ctx = {
        "request": request,
        "username": user["username"],
        "role": user["role"],
        "csrf_token": _csrf_token(request),
        "snap": snap,
        "db_ok": obs_store.db_available(),
        "mode": snap.get("mode", "REAL"),
        "system_state": _system_state(snap),
        "alert_cls": alert_cls,
        "alert_text": alert_text,
        "led": led,
        "commands_enabled": "YES" if COMMANDS_ENABLED else "NO",
    }
    if extra:
        ctx.update(extra)
    return ctx


def _system_state(snap):
    if snap.get("mode") == "SIMULATION":
        return "◆ SIMULATION"
    st = snap.get("telemetry_status", "UNKNOWN")
    if st == "LIVE":
        return "● READY"
    if st == "STALE":
        return "⚠ STALE"
    return "○ DISCONNECTED"


def _hms(hours):
    if hours is None:
        return "—"
    hours = hours % 24.0
    h = int(hours)
    m = int((hours - h) * 60.0)
    s = (hours - h - m / 60.0) * 3600.0
    return f"{h:02d}:{m:02d}:{s:04.1f}"


# ── Workspace pages ─────────────────────────────────────────────────────
@router.get("/app/control", include_in_schema=False)
async def page_control(request: Request):
    user = auth.get_current_user(request)
    if user is None:
        return RedirectResponse(url="/api/login", status_code=302)
    snap = obs_svc.observatory_snapshot()

    # Use rich telemetry from unified snapshot (prefers WS, falls back to Alpaca)
    ra = snap.get("ra_hours")
    dec = snap.get("dec_deg")
    ha = snap.get("ha_hours")
    lst = snap.get("lst_hours")
    alt = snap.get("alt_deg")
    az = snap.get("az_deg")
    dome_az = snap.get("dome_az_deg")
    tracking = snap.get("tracking")
    slewing = snap.get("slewing")
    dome_state = snap.get("dome_state")
    ra_speed = snap.get("ra_speed")
    ra_direction = snap.get("ra_direction")
    dec_speed = snap.get("dec_speed")
    dec_direction = snap.get("dec_direction")
    safety_limit_active = snap.get("safety_limit_active")
    safety_limit_message = snap.get("safety_limit_message")
    cooldown_active = snap.get("cooldown_active")
    telemetry_source = snap.get("telemetry_source", "NONE")

    # HRA from WebSocket if available, else compute from LST - RA
    hra = ha if ha is not None else (((lst - ra) % 24.0) if (lst is not None and ra is not None) else None)

    import time as _t
    # Age based on the active telemetry source
    ws_age = _t.time() - (snap.get("ws", {}).get("updated_at") or 0.0)
    alpaca_age = _t.time() - (snap.get("alpaca", {}).get("updated_at") or 0.0)
    telemetry_age = ws_age if telemetry_source == "WS" else alpaca_age

    wvals = snap["weather"].get("values") or {}
    wx_age = (f"{snap['weather'].get('age_s', 0.0):.1f}s" if wvals else "—")

    # Motion state from rich telemetry
    if slewing is True:
        motion_state = "SLEWING"
    elif slewing is False:
        motion_state = "IDLE"
    else:
        motion_state = "UNKNOWN"

    # Dome state from WebSocket
    if dome_state:
        dome_state_str = dome_state
    elif slewing is True:
        dome_state_str = "MOVING"
    elif slewing is False:
        dome_state_str = "IDLE"
    else:
        dome_state_str = "UNKNOWN"

    ctx = _page_ctx(request, user, snap, {
        "page_title": "Control", "active": "control",
        "ra_hms": obs_svc.fmt_ra_hms(ra),
        "hra_hms": obs_svc.fmt_ra_hms(hra),
        "dec_dms": obs_svc.fmt_dec_dms(dec),
        "az": obs_svc.fmt_deg(az),
        "alt": obs_svc.fmt_deg(alt),
        "lst_hms": _hms(lst),
        "jd": f"{snap['jd']:.5f}",
        "dome_az": obs_svc.fmt_deg(dome_az) if dome_az is not None else "—",
        "dome_state": dome_state_str,
        "dome_sync": "Unknown",  # Would need additional telemetry
        "tracking": ("ON" if tracking is True
                      else ("OFF" if tracking is False else "UNKNOWN")),
        "sidereal": ("ON" if tracking is True
                      else ("OFF" if tracking is False else "Unknown")),
        "slewing": slewing,
        "atpark": snap.get("alpaca", {}).get("values", {}).get("atpark"),
        "athome": snap.get("alpaca", {}).get("values", {}).get("athome"),
        "motion_state": motion_state,
        "ra_speed": ra_speed,
        "ra_direction": ra_direction,
        "dec_speed": dec_speed,
        "dec_direction": dec_direction,
        "safety_limit_active": safety_limit_active,
        "safety_limit_message": safety_limit_message,
        "cooldown_active": cooldown_active,
        "telemetry_age": f"{telemetry_age:.1f}s",
        "telemetry_source": telemetry_source,
        "wx_age": wx_age,
        "dec_home": ((snap.get("config") or {}).get("state") or {}).get("dec_home_deg", "—"),
        "ra_off": ((snap.get("config") or {}).get("state") or {}).get("minor_hra_adj", "—"),
        "dec_off": ((snap.get("config") or {}).get("state") or {}).get("minor_dec_adj", "—"),
        "dome_off": ((snap.get("config") or {}).get("state") or {}).get("dome_offset", "—"),
        "site_lat": snap.get("alpaca", {}).get("values", {}).get("sitelatitude") or "29.36",
        "site_lon": snap.get("alpaca", {}).get("values", {}).get("sitelongitude") or "79.45",
        "cmd_result": request.query_params.get("cmd", ""),
        # Environment/safety breakdown
        "env": snap.get("environment", {}),
    })
    return template_response("control.html", **ctx)


@router.get("/app/checkup", include_in_schema=False)
async def page_checkup_retired(request: Request):
    """CHECKUP was merged into CONTROL — redirect, never a dead link."""
    return RedirectResponse(url="/app/control", status_code=302)


@router.get("/app/history", include_in_schema=False)
async def page_history_retired(request: Request):
    """Old history URL → canonical observations archive."""
    return RedirectResponse(url="/app/observations", status_code=302)


@router.get("/app/syscam", include_in_schema=False)
async def page_syscam_retired(request: Request):
    """Retired merged page → SYSTEM (camera lives at /app/camera)."""
    return RedirectResponse(url="/app/system", status_code=302)


@router.get("/app/observations", include_in_schema=False)
async def page_history(request: Request):
    user = auth.get_current_user(request)
    if user is None:
        return RedirectResponse(url="/api/login", status_code=302)
    snap = obs_svc.observatory_snapshot()
    qp = request.query_params
    page = max(1, int(qp.get("page", "1") or 1))
    limit = 15
    filt = {"q": qp.get("q", ""), "status": qp.get("status", ""),
            "date_from": qp.get("from", ""), "date_to": qp.get("to", "")}
    owner_filter = qp.get("owner", "")
    eff_user, eff_role = user["username"], user["role"]
    if user["role"] == "admin" and owner_filter:
        eff_user, eff_role = owner_filter, "scientist"
    try:
        res = obs_store.list_observations(eff_user, eff_role, limit=limit,
                                          offset=(page - 1) * limit, **filt)
        try:
            mine = obs_store.list_observations(user["username"], user["role"],
                                               limit=1000)
            durations = [m.get("duration_s") or 0 for m in mine["items"]]
            summary = {"total": mine["total"],
                       "hours": round(sum(durations) / 3600.0, 2),
                       "targets": len({m.get("target") for m in mine["items"]}),
                       "latest": mine["items"][0]["start"] if mine["items"] else "—"}
        except Exception:
            summary = {"total": res["total"], "hours": "—", "targets": "—", "latest": "—"}
    except Exception:
        res = {"total": 0, "items": []}
        summary = {"total": 0, "hours": "—", "targets": "—", "latest": "—"}
    detail = None
    if qp.get("obs"):
        try:
            detail = obs_store.get_observation_with_files(int(qp["obs"]), user["username"], user["role"])
        except Exception:
            detail = None
    from datetime import datetime as _dt, timedelta as _td, timezone as _tz
    _now = _dt.now(_tz.utc).date()
    ctx = _page_ctx(request, user, snap, {
        "page_title": "Observations", "active": "observations",
        "items": res["items"], "total": res["total"], "page": page,
        "summary": summary, "q": filt["q"], "status_f": filt["status"],
        "date_from": filt["date_from"], "date_to": filt["date_to"],
        "owner_filter": owner_filter, "detail": detail,
        "today": _now.isoformat(),
        "week_ago": (_now - _td(days=7)).isoformat(),
        "month_ago": (_now - _td(days=30)).isoformat(),
    })
    return template_response("observations.html", **ctx)


@router.get("/app/camera", include_in_schema=False)
async def page_camera(request: Request):
    """Dedicated camera workspace: status, acquisition, images."""
    user = auth.get_current_user(request)
    if user is None:
        return RedirectResponse(url="/api/login", status_code=302)
    snap = obs_svc.observatory_snapshot()
    # Get camera status from the camera service
    camera_status = {}
    try:
        from .camera_service import get_camera_service
        svc = get_camera_service()
        if svc:
            camera_status = svc.get_camera_status()
    except Exception:
        camera_status = {"available": False, "state": "NOT_AVAILABLE"}
    return template_response("camera.html", **_page_ctx(
        request, user, snap, {"page_title": "Camera", "active": "camera",
                               "camera_status": camera_status}))


@router.get("/app/system", include_in_schema=False)
async def page_system(request: Request):
    """Engineering diagnostics: health, connections, logs, configuration."""
    user = auth.get_current_user(request)
    if user is None:
        return RedirectResponse(url="/api/login", status_code=302)
    snap = obs_svc.observatory_snapshot()
    cfg = snap.get("config") or {}
    limits = (cfg.get("limits") or {}).get("motion", {})
    audit = _recent_audit()
    db_ok = obs_store.db_available()
    # WHY rewrite two rows: the snapshot builder cannot see PostgreSQL, so it
    # reports Database/WebApp as UNKNOWN. Here the real states are known.
    subs = []
    for name, st, ts in snap.get("subsystems", []):
        if name == "Database":
            subs.append((name, "LIVE" if db_ok else "NOT_CONNECTED", ts))
        elif name == "Web Application":
            import time as _t3
            subs.append((name, "LIVE", _t3.time()))
        else:
            subs.append((name, st, ts))
    snap = dict(snap, subsystems=subs)
    ctx = _page_ctx(request, user, snap, {
        "page_title": "System", "active": "system", "audit": audit,
        "alt_floor": limits.get("min_altitude_deg", "—"),
        "alt_max": limits.get("max_altitude_deg", "—"),
        "site_lat": "29.36", "site_lon": "79.45",
        "env": snap.get("environment", {}),
        "env_thresholds": (cfg.get("limits") or {}).get("environment", {}),
    })
    return template_response("system.html", **ctx)


@router.get("/app/admin", include_in_schema=False)
async def page_admin(request: Request):
    user = auth.get_current_user(request)
    if user is None or user["role"] != "admin":
        return RedirectResponse(url="/api/login", status_code=302)
    snap = obs_svc.observatory_snapshot()
    users, items = [], []
    try:
        from .database import get_db_connection
        conn = get_db_connection()
        try:
            with conn.cursor() as cur:
                cur.execute("SELECT username, role, is_active FROM users ORDER BY username")
                users = [{"username": r[0], "role": r[1], "is_active": r[2]}
                         for r in cur.fetchall()]
        finally:
            conn.close()
    except Exception:
        users = [{"username": user["username"], "role": user["role"], "is_active": True}]
    try:
        items = obs_store.list_observations(user["username"], "admin", limit=50)["items"]
    except Exception:
        items = []
    return template_response("admin.html", **_page_ctx(
        request, user, snap, {"page_title": "Admin", "active": "admin",
                              "users": users, "items": items,
                              "audit": _recent_audit()}))


def _recent_audit(limit=50):
    try:
        from .database import get_db_connection
        conn = get_db_connection()
        try:
            with conn.cursor() as cur:
                cur.execute("SELECT actor_username, action, detail, created_at "
                            "FROM audit_events ORDER BY id DESC LIMIT %s", (limit,))
                return cur.fetchall()
        finally:
            conn.close()
    except Exception:
        return []


# ── Read-only JSON APIs ─────────────────────────────────────────────────
@router.get("/api/telemetry/snapshot", include_in_schema=False)
async def api_snapshot(request: Request):
    _require_user(request)
    snap = obs_svc.observatory_snapshot()
    return {
        "mode": snap["mode"],
        "telemetry_status": snap.get("telemetry_status"),
        "telemetry_source": snap.get("telemetry_source"),
        "ra_hours": snap.get("ra_hours"),
        "dec_deg": snap.get("dec_deg"),
        "ha_hours": snap.get("ha_hours"),
        "lst_hours": snap.get("lst_hours"),
        "lst_source": snap.get("lst_source"),
        "alt_deg": snap.get("alt_deg"),
        "az_deg": snap.get("az_deg"),
        "dome_az_deg": snap.get("dome_az_deg"),
        "tracking": snap.get("tracking"),
        "slewing": snap.get("slewing"),
        "ra_speed": snap.get("ra_speed"),
        "ra_direction": snap.get("ra_direction"),
        "dec_speed": snap.get("dec_speed"),
        "dec_direction": snap.get("dec_direction"),
        "dome_state": snap.get("dome_state"),
        "safety_limit_active": snap.get("safety_limit_active"),
        "safety_limit_message": snap.get("safety_limit_message"),
        "cooldown_active": snap.get("cooldown_active"),
        "atpark": snap.get("alpaca", {}).get("values", {}).get("atpark"),
        "athome": snap.get("alpaca", {}).get("values", {}).get("athome"),
        "weather": snap["weather"],
        "safety": snap["safety"],
        "subsystems": snap["subsystems"],
        "generated_at": snap["generated_at"],
        "environment": snap.get("environment", {}),
    }


@router.get("/api/history/list", include_in_schema=False)
async def api_history_list(request: Request):
    user = _require_user(request)
    qp = request.query_params
    try:
        return obs_store.list_observations(
            user["username"], user["role"], q=qp.get("q", ""),
            status=qp.get("status", ""), date_from=qp.get("from", ""),
            date_to=qp.get("to", ""),
            limit=min(100, int(qp.get("limit", "25") or 25)),
            offset=int(qp.get("offset", "0") or 0))
    except Exception as e:
        raise HTTPException(status_code=503, detail=f"Database unavailable: {e}")


@router.get("/api/history/export", include_in_schema=False)
async def api_history_export(request: Request, fmt: str = "csv"):
    """Export ONLY records the caller is authorized to see."""
    import logging
    _log = logging.getLogger("Export")
    user = _require_user(request)
    try:
        res = obs_store.list_observations(user["username"], user["role"], limit=5000)
    except Exception as e:
        _log.exception("Export query failed")
        raise HTTPException(status_code=503, detail=f"Database unavailable: {e}")
    obs_store.record_audit(user["username"], "observation_export",
                           f"fmt={fmt} count={res['total']}")
    if fmt == "json":
        return JSONResponse(res)
    cols = ["id", "owner", "target", "ra_deg", "dec_deg", "start",
            "end", "duration_s", "status"]
    rows = [[o["id"], o["owner"], o["target"], o["ra_deg"], o["dec_deg"],
             o["start"], o["end"], o["duration_s"], o["status"]]
            for o in res["items"]]
    if fmt == "xlsx":
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Font
        except ImportError:
            _log.error("openpyxl not installed")
            raise HTTPException(status_code=500, detail="Excel export unavailable (openpyxl not installed)")
        try:
            wb = Workbook()
            ws = wb.active
            ws.title = "Observations"
            ws.append(cols)
            for r in rows:
                ws.append([str(v) if v is not None else "" for v in r])
            bold = Font(bold=True)
            for cell in ws[1]:
                cell.font = bold
            buf = io.BytesIO()
            wb.save(buf)
            buf.seek(0)
            content = buf.getvalue()
            return StreamingResponse(
                iter([content]),
                media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                headers={"Content-Disposition": "attachment; filename=observations.xlsx"}
            )
        except Exception as e:
            _log.exception("XLSX export failed")
            raise HTTPException(status_code=500, detail=f"Excel export failed: {e}")
    if fmt == "pdf":
        try:
            from fpdf import FPDF
        except ImportError:
            _log.error("fpdf not installed")
            raise HTTPException(status_code=500, detail="PDF export unavailable (fpdf not installed)")
        try:
            pdf = FPDF(orientation="L", format="A4")
            pdf.set_auto_page_break(auto=True, margin=15)
            pdf.add_page()
            pdf.set_font("Helvetica", "B", 14)
            pdf.cell(0, 10, "ARIES 104 cm Sampurnanand Telescope - Observation Export",
                     new_x="LMARGIN", new_y="NEXT")
            pdf.set_font("Helvetica", "", 9)
            pdf.cell(0, 6,
                     f"Exported by {user['username']} ({user['role']}) - {len(rows)} record(s), privacy-filtered",
                     new_x="LMARGIN", new_y="NEXT")
            pdf.ln(2)
            widths = [12, 30, 60, 22, 22, 38, 38, 22, 28]
            pdf.set_font("Helvetica", "B", 8)
            for i, c in enumerate(cols):
                pdf.cell(widths[i], 7, c, border=1)
            pdf.ln()
            pdf.set_font("Helvetica", "", 8)
            for r in rows:
                for i, v in enumerate(r):
                    txt = str(v) if v is not None else "-"
                    pdf.cell(widths[i], 6, txt, border=1)
                pdf.ln()
            out = pdf.output()
            if isinstance(out, str):
                out = out.encode("latin-1", errors="replace")
            elif not isinstance(out, bytes):
                out = bytes(out)
            return StreamingResponse(
                iter([out]),
                media_type="application/pdf",
                headers={"Content-Disposition": "attachment; filename=observations.pdf"}
            )
        except Exception as e:
            _log.exception("PDF export failed")
            raise HTTPException(status_code=500, detail=f"PDF export failed: {e}")
    # Default: CSV
    try:
        buf = io.StringIO()
        w = csv.writer(buf)
        w.writerow(cols)
        w.writerows(rows)
        content = buf.getvalue()
        return StreamingResponse(
            iter([content.encode("utf-8")]),
            media_type="text/csv; charset=utf-8",
            headers={"Content-Disposition": "attachment; filename=observations.csv"}
        )
    except Exception as e:
        _log.exception("CSV export failed")
        raise HTTPException(status_code=500, detail=f"CSV export failed: {e}")


# ── Safe command gate (validate first, execute only when enabled) ───────
def _validate_command(cmd, fields):
    """Mirror the existing STCS validation sets. Returns (ok, message)."""
    if cmd not in VALID_COMMANDS:
        return False, f"Unknown command '{cmd}'"
    if cmd == "manual_move":
        axis = (fields.get("axis") or "").upper()
        direction = (fields.get("direction") or "").upper()
        speed = (fields.get("speed") or "").upper()
        if axis not in VALID_DIRECTIONS:
            return False, "axis must be RA or DEC"
        if direction not in VALID_DIRECTIONS[axis]:
            return False, f"direction '{direction}' invalid for {axis}"
        if speed not in VALID_SPEEDS:
            return False, f"speed must be one of {sorted(VALID_SPEEDS)}"
        return True, "manual_move valid"
    if cmd == "slewtocoordinatesasync":
        try:
            ra = float(fields.get("ra_deg", ""))
            dec = float(fields.get("dec_deg", ""))
        except (TypeError, ValueError):
            return False, "ra_deg/dec_deg must be numbers"
        if not (-90.0 <= dec <= 90.0):
            return False, "dec out of range"
        return True, f"slew target RA={ra} DEC={dec} valid"
    return True, f"{cmd} valid"


@router.post("/api/command/validate", include_in_schema=False)
async def api_command_validate(request: Request):
    """Validate a command intent. Executes NOTHING unless explicitly enabled.

    Even when STCS_COMMANDS_ENABLED=1, execution requires the caller to hold
    the control lock AND a reviewed hardware bridge (not present in Phase 2),
    so this endpoint returns executed:false with the reason. Physical motion
    stays reserved for controlled validation (RULE 20).
    """
    user = _require_user(request)
    form = await request.form()
    if not _csrf_validate_form(form, request):
        raise HTTPException(status_code=403, detail="Invalid CSRF token")
    fields = dict(form)
    cmd = (fields.get("command") or "").strip()
    ok, message = _validate_command(cmd, fields)
    if not ok:
        obs_store.record_audit(user["username"], "command_rejected", f"{cmd}: {message}")
        raise HTTPException(status_code=400, detail=message)
    lock = obs_svc.control_lock_status()

    def _respond(payload):
        """Browsers (form posts) go back to CONTROL; API clients keep JSON."""
        accept = request.headers.get("accept", "")
        if "text/html" in accept and payload.get("redirect"):
            return RedirectResponse(url=payload["redirect"], status_code=302)
        return payload

    if not COMMANDS_ENABLED:
        obs_store.record_audit(user["username"], "command_validated", f"{cmd}: {message} (not executed; commands disabled)")
        return _respond({"ok": True, "executed": False, "command": cmd, "message": message,
                "reason": "STCS_COMMANDS_ENABLED=0 — validation only, no motion",
                "redirect": f"/app/control?cmd={cmd}+validated+(no+motion)"})
    if lock["owner"] != user["username"]:
        return _respond({"ok": False, "executed": False, "command": cmd,
                "reason": f"Control lock held by {lock['owner'] or 'nobody'} — acquire it first",
                "redirect": "/app/control?cmd=lock+required"})
    obs_store.record_audit(user["username"], "command_validated", f"{cmd}: {message} (bridge pending review)")
    return _respond({"ok": True, "executed": False, "command": cmd, "message": message,
            "reason": "Hardware bridge pending RULE-20 review — no motion executed",
            "redirect": "/app/control?cmd=bridge+pending+review"})


@router.post("/api/command/lock", include_in_schema=False)
async def api_lock(request: Request):
    user = _require_user(request)
    form = await request.form()
    if not _csrf_validate_form(form, request):
        raise HTTPException(status_code=403, detail="Invalid CSRF token")
    res = obs_svc.acquire_control_lock(user["username"])
    obs_store.record_audit(user["username"], "control_lock",
                           f"acquire ok={res['ok']} owner={res.get('owner')}")
    return RedirectResponse(url="/app/control", status_code=302)


@router.post("/api/command/unlock", include_in_schema=False)
async def api_unlock(request: Request):
    user = _require_user(request)
    form = await request.form()
    if not _csrf_validate_form(form, request):
        raise HTTPException(status_code=403, detail="Invalid CSRF token")
    obs_svc.release_control_lock(user["username"], user["role"])
    obs_store.record_audit(user["username"], "control_lock", "release")
    return RedirectResponse(url="/app/control", status_code=302)


# ── Observation Archive APIs ─────────────────────────────────────────────

@router.post("/api/observations/delete/{obs_id}", include_in_schema=False)
async def api_delete_observation(request: Request, obs_id: int):
    """Delete an observation. Scientists can only delete their own."""
    user = _require_user(request)
    form = await request.form()
    if not _csrf_validate_form(form, request):
        raise HTTPException(status_code=403, detail="Invalid CSRF token")
    deleted, reason = obs_store.delete_observation(obs_id, user["username"], user["role"])
    obs_store.record_audit(user["username"], "observation_delete",
                           f"obs_id={obs_id} deleted={deleted} reason={reason}")
    if not deleted:
        raise HTTPException(status_code=403, detail=reason)
    return RedirectResponse(url="/app/observations", status_code=302)


@router.get("/api/observations/summary", include_in_schema=False)
async def api_observations_summary(request: Request):
    """Get observation summary statistics."""
    user = _require_user(request)
    try:
        return obs_store.get_observation_summary(user["username"], user["role"])
    except Exception as e:
        raise HTTPException(status_code=503, detail=f"Database unavailable: {e}")


@router.get("/api/observations/scientist/{scientist}", include_in_schema=False)
async def api_observations_scientist(request: Request, scientist: str):
    """Get observations for a specific scientist (admin only)."""
    user = _require_user(request)
    if user["role"] != "admin":
        raise HTTPException(status_code=403, detail="Admin access required")
    qp = request.query_params
    try:
        return obs_store.list_observations_for_scientist(
            scientist,
            limit=min(100, int(qp.get("limit", "25") or 25)),
            offset=int(qp.get("offset", "0") or 0)
        )
    except Exception as e:
        raise HTTPException(status_code=503, detail=f"Database unavailable: {e}")


@router.get("/api/admin/scientists", include_in_schema=False)
async def api_admin_scientists(request: Request):
    """List all scientists with observation counts (admin only)."""
    user = _require_user(request)
    if user["role"] != "admin":
        raise HTTPException(status_code=403, detail="Admin access required")
    try:
        return obs_store.list_scientists()
    except Exception as e:
        raise HTTPException(status_code=503, detail=f"Database unavailable: {e}")


@router.get("/api/admin/audit", include_in_schema=False)
async def api_admin_audit(request: Request):
    """Get audit events (admin only)."""
    user = _require_user(request)
    if user["role"] != "admin":
        raise HTTPException(status_code=403, detail="Admin access required")
    qp = request.query_params
    try:
        return obs_store.get_audit_events(
            limit=min(200, int(qp.get("limit", "50") or 50)),
            offset=int(qp.get("offset", "0") or 0),
            action_filter=qp.get("action", "")
        )
    except Exception as e:
        raise HTTPException(status_code=503, detail=f"Database unavailable: {e}")
